from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from invoice_agent.exporters import export_csv, export_xlsx  # noqa: E402
from invoice_agent.extractor import extract_invoice  # noqa: E402


class ExportersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        pdf_path = ROOT / "output" / "pdf" / "dea_test_invoice.pdf"
        cls.invoice = extract_invoice(pdf_path).to_dict()
        cls.temp_root = ROOT / "tmp"
        cls.temp_root.mkdir(parents=True, exist_ok=True)

    def test_csv_export(self) -> None:
        with tempfile.TemporaryDirectory(dir=self.temp_root) as directory:
            output = Path(directory) / "register.csv"
            export_csv([self.invoice], output)
            with output.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.DictReader(stream, delimiter=";"))
            self.assertEqual(1, len(rows))
            self.assertEqual("DEA-1042", rows[0]["Номер счёта"])
            self.assertEqual("2400480.0", rows[0]["Всего к оплате"])

    def test_xlsx_export(self) -> None:
        with tempfile.TemporaryDirectory(dir=self.temp_root) as directory:
            output = Path(directory) / "register.xlsx"
            export_xlsx([self.invoice], output)
            workbook = load_workbook(output, data_only=True)
            sheet = workbook["Реестр счетов"]
            headers = [cell.value for cell in sheet[1]]
            values = [cell.value for cell in sheet[2]]
            row = dict(zip(headers, values))
            self.assertEqual("DEA-1042", row["Номер счёта"])
            self.assertEqual(2400480.0, row["Всего к оплате"])
            self.assertEqual(6, row["Позиций"])

            item_sheet = workbook["Товарные позиции"]
            self.assertEqual(7, item_sheet.max_row)
            item_headers = [cell.value for cell in item_sheet[1]]
            first_values = [cell.value for cell in item_sheet[2]]
            first_item = dict(zip(item_headers, first_values))
            self.assertEqual("Беговая дорожка D.E.A. RUN T4", first_item["Наименование"])
            self.assertEqual(708000.0, first_item["Сумма"])


if __name__ == "__main__":
    unittest.main()

