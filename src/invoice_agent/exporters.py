from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


COLUMNS: list[tuple[str, str]] = [
    ("source_file", "Файл"),
    ("supplier_name", "Поставщик"),
    ("supplier_inn", "ИНН поставщика"),
    ("supplier_kpp", "КПП поставщика"),
    ("buyer_name", "Покупатель"),
    ("buyer_inn", "ИНН покупателя"),
    ("invoice_number", "Номер счёта"),
    ("invoice_date", "Дата счёта"),
    ("subtotal_amount", "Сумма без НДС"),
    ("vat_rate", "НДС, %"),
    ("vat_amount", "Сумма НДС"),
    ("total_amount", "Всего к оплате"),
    ("currency", "Валюта"),
    ("payment_deadline", "Срок оплаты"),
    ("contract_number", "Договор"),
    ("item_count", "Позиций"),
    ("warnings", "Предупреждения"),
]


def flatten_invoice(invoice: dict[str, Any]) -> dict[str, Any]:
    supplier = invoice.get("supplier") or {}
    buyer = invoice.get("buyer") or {}
    return {
        "source_file": invoice.get("source_file"),
        "supplier_name": supplier.get("name"),
        "supplier_inn": supplier.get("inn"),
        "supplier_kpp": supplier.get("kpp"),
        "buyer_name": buyer.get("name"),
        "buyer_inn": buyer.get("inn"),
        "invoice_number": invoice.get("invoice_number"),
        "invoice_date": invoice.get("invoice_date"),
        "subtotal_amount": invoice.get("subtotal_amount"),
        "vat_rate": invoice.get("vat_rate"),
        "vat_amount": invoice.get("vat_amount"),
        "total_amount": invoice.get("total_amount"),
        "currency": invoice.get("currency"),
        "payment_deadline": invoice.get("payment_deadline"),
        "contract_number": invoice.get("contract_number"),
        "item_count": len(invoice.get("items") or []),
        "warnings": "; ".join(invoice.get("warnings") or []),
    }


def export_csv(invoices: Iterable[dict[str, Any]], output: str | Path) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows = [flatten_invoice(invoice) for invoice in invoices]
    keys = [key for key, _ in COLUMNS]
    headers = {key: label for key, label in COLUMNS}

    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
    return output_path


def export_xlsx(invoices: Iterable[dict[str, Any]], output: str | Path) -> Path:
    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    invoice_list = list(invoices)
    rows = [flatten_invoice(invoice) for invoice in invoice_list]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реестр счетов"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    header_fill = PatternFill("solid", fgColor="0F766E")
    warning_fill = PatternFill("solid", fgColor="FEF3C7")
    for column_index, (_, label) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _) in enumerate(COLUMNS, start=1):
            value = row.get(key)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key == "warnings" and value:
                cell.fill = warning_fill

    widths = {
        "A": 28,
        "B": 34,
        "C": 17,
        "D": 16,
        "E": 30,
        "F": 17,
        "G": 17,
        "H": 14,
        "I": 18,
        "J": 11,
        "K": 16,
        "L": 18,
        "M": 10,
        "N": 15,
        "O": 18,
        "P": 11,
        "Q": 48,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.row_dimensions[1].height = 28
    sheet.sheet_view.showGridLines = False

    item_sheet = workbook.create_sheet("Товарные позиции")
    item_columns = [
        "Файл",
        "Номер счёта",
        "№ позиции",
        "Наименование",
        "Единица",
        "Количество",
        "Цена",
        "Сумма",
        "Валюта",
    ]
    for column_index, label in enumerate(item_columns, start=1):
        cell = item_sheet.cell(row=1, column=column_index, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    item_row = 2
    for invoice in invoice_list:
        for item in invoice.get("items") or []:
            values = [
                invoice.get("source_file"),
                invoice.get("invoice_number"),
                item.get("line_number"),
                item.get("name"),
                item.get("unit"),
                item.get("quantity"),
                item.get("unit_price"),
                item.get("line_total"),
                invoice.get("currency"),
            ]
            for column_index, value in enumerate(values, start=1):
                cell = item_sheet.cell(row=item_row, column=column_index, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            item_row += 1

    for column, width in {
        "A": 28,
        "B": 17,
        "C": 12,
        "D": 48,
        "E": 13,
        "F": 14,
        "G": 16,
        "H": 16,
        "I": 10,
    }.items():
        item_sheet.column_dimensions[column].width = width
    item_sheet.freeze_panes = "A2"
    item_sheet.auto_filter.ref = f"A1:I{max(1, item_row - 1)}"
    item_sheet.sheet_view.showGridLines = False
    item_sheet.row_dimensions[1].height = 28
    workbook.save(output_path)
    return output_path

